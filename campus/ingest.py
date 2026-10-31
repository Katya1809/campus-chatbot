"""Multi-source ingestion: PDF, DOCX, MD/TXT, CSV, ICS calendars and web pages -> LangChain Documents."""
from __future__ import annotations

import re
import urllib.robotparser as rp
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import requests
from bs4 import BeautifulSoup
from langchain_core.documents import Document

CATEGORIES = {
    "admissions": ["admission", "eligibility", "verification", "fee", "scholarship", "tuition"],
    "academics": ["semester", "cbcs", "credit", "exam", "internal", "external", "attendance",
                  "syllabus", "grade", "cgpa", "backlog", "timetable"],
    "facilities": ["library", "hostel", "canteen", "mess", "sports", "gym", "medical", "lab",
                   "auditorium", "parking", "block"],
    "clubs": ["club", "society", "drama", "music", "dance", "nss", "ncc", "committee", "chapter"],
    "events": ["event", "fest", "workshop", "hackathon", "seminar", "competition", "cultural"],
    "placement": ["placement", "internship", "recruiter", "company", "drive", "training"],
    "transport": ["bus", "shuttle", "transport", "route"],
    "emergency": ["emergency", "security", "ambulance", "helpline", "ragging", "grievance"],
    "digital": ["portal", "erp", "lms", "login", "wifi", "online"],
    "contacts": ["contact", "phone", "email", "hod", "dean", "principal", "faculty"],
}
CSV_TYPES = {"events": "events", "contacts": "contacts", "clubs": "clubs", "locations": "facilities"}

# Minimum characters a text chunk must have to be indexed
_MIN_CHUNK = 50


def categorize(text: str) -> str:
    words = re.findall(r"[a-z0-9]+", text.lower())
    scores = {c: sum(words.count(k) for k in kws) for c, kws in CATEGORIES.items()}
    best = max(scores, key=scores.get)
    return best if scores[best] else "general"


def _meta(source: str, kind: str, category: str) -> dict:
    return {"source": source, "type": kind, "category": category,
            "updated": datetime.now().strftime("%Y-%m-%d %H:%M")}


# ---------------------------------------------------------------------------
# PDF loader — page-by-page with fallback section splitting
# ---------------------------------------------------------------------------

def load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader
    docs = []
    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if not text or len(text) < _MIN_CHUNK:
            continue
        # Split on double newlines to get finer sections within a page
        sections = [s.strip() for s in re.split(r"\n{2,}", text) if len(s.strip()) >= _MIN_CHUNK]
        if not sections:
            sections = [text]
        for sec in sections:
            m = _meta(f"{path.name} (p.{i})", "handbook", categorize(sec))
            docs.append(Document(page_content=sec, metadata=m))
    return docs


# ---------------------------------------------------------------------------
# DOCX loader
# ---------------------------------------------------------------------------

def load_docx(path: Path) -> list[Document]:
    import docx
    d = docx.Document(str(path))
    # Group paragraphs into logical sections split by heading styles
    sections: list[list[str]] = [[]]
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        if para.style.name.startswith("Heading"):
            if sections[-1]:
                sections.append([])
        sections[-1].append(t)
    docs = []
    for sec in sections:
        text = "\n".join(sec).strip()
        if len(text) >= _MIN_CHUNK:
            docs.append(Document(page_content=text, metadata=_meta(path.name, "handbook", categorize(text))))
    return docs or [Document(page_content="(empty)", metadata=_meta(path.name, "handbook", "general"))]


# ---------------------------------------------------------------------------
# Markdown / plain-text loader — split by headings
# ---------------------------------------------------------------------------

def load_text(path: Path) -> list[Document]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    parts = re.split(r"\n(?=#{1,3} )", raw)
    return [Document(page_content=p.strip(),
                     metadata=_meta(path.name, "handbook", categorize(p)))
            for p in parts if len(p.strip()) >= _MIN_CHUNK]


# ---------------------------------------------------------------------------
# CSV loader
# ---------------------------------------------------------------------------

def load_csv(path: Path) -> list[Document]:
    df = pd.read_csv(path).fillna("")
    kind = CSV_TYPES.get(path.stem)
    docs = []
    for _, row in df.iterrows():
        text = " | ".join(f"{c}: {v}" for c, v in row.items() if str(v).strip())
        if len(text) < _MIN_CHUNK:
            continue
        cat = kind or categorize(text)
        docs.append(Document(page_content=text, metadata=_meta(path.name, kind or "csv", cat)))
    return docs


# ---------------------------------------------------------------------------
# ICS (iCalendar) loader — NEW
# ---------------------------------------------------------------------------

def load_ics(path: Path) -> list[Document]:
    """Parse an iCalendar (.ics) file into one Document per VEVENT."""
    try:
        from icalendar import Calendar  # type: ignore
    except ImportError:
        # Fallback: simple line-by-line parser
        return _load_ics_fallback(path)

    docs = []
    try:
        cal = Calendar.from_ical(path.read_bytes())
        for component in cal.walk():
            if component.name != "VEVENT":
                continue
            summary = str(component.get("SUMMARY", "Untitled Event"))
            dtstart = component.get("DTSTART")
            location = str(component.get("LOCATION", ""))
            description = str(component.get("DESCRIPTION", ""))
            categories = str(component.get("CATEGORIES", ""))
            organizer = str(component.get("ORGANIZER", ""))
            date_str = ""
            if dtstart:
                try:
                    date_str = dtstart.dt.strftime("%Y-%m-%d %H:%M") if hasattr(dtstart.dt, "strftime") else str(dtstart.dt)
                except Exception:
                    date_str = str(dtstart)
            text = (f"Event: {summary}\nDate: {date_str}\nLocation: {location}\n"
                    f"Description: {description}\nCategories: {categories}\nOrganizer: {organizer}")
            text = text.strip()
            if len(text) >= _MIN_CHUNK:
                m = _meta(path.name, "ics", categorize(text))
                docs.append(Document(page_content=text, metadata=m))
    except Exception:
        return _load_ics_fallback(path)
    return docs


def _load_ics_fallback(path: Path) -> list[Document]:
    """Simple regex-based ICS parser when icalendar package is unavailable."""
    raw = path.read_text(encoding="utf-8", errors="ignore")
    events = re.split(r"BEGIN:VEVENT", raw)[1:]
    docs = []
    for ev in events:
        summary = re.search(r"SUMMARY:(.*)", ev)
        dtstart = re.search(r"DTSTART[^:]*:(.*)", ev)
        location = re.search(r"LOCATION:(.*)", ev)
        description = re.search(r"DESCRIPTION:(.*?)(?:\r?\n[A-Z])", ev, re.DOTALL)
        parts = []
        if summary:
            parts.append(f"Event: {summary.group(1).strip()}")
        if dtstart:
            parts.append(f"Date: {dtstart.group(1).strip()}")
        if location:
            parts.append(f"Location: {location.group(1).strip()}")
        if description:
            parts.append(f"Description: {description.group(1).strip()[:300]}")
        text = "\n".join(parts)
        if len(text) >= _MIN_CHUNK:
            m = _meta(path.name, "ics", "events")
            docs.append(Document(page_content=text, metadata=m))
    return docs


# ---------------------------------------------------------------------------
# Web scraper
# ---------------------------------------------------------------------------

def _allowed(url: str) -> bool:
    try:
        u = urlparse(url)
        parser = rp.RobotFileParser(f"{u.scheme}://{u.netloc}/robots.txt")
        parser.read()
        return parser.can_fetch("CampusInfoBot", url)
    except Exception:
        return True  # robots.txt unreachable -> proceed politely


def _extract_structured(soup: BeautifulSoup) -> dict[str, list[str]]:
    """Pull tables, lists and headings out of a page for richer context."""
    structured: dict[str, list[str]] = {}

    # Tables -> each row as a string
    tables: list[str] = []
    for table in soup.find_all("table"):
        rows = []
        for tr in table.find_all("tr"):
            cells = [td.get_text(" ", strip=True) for td in tr.find_all(["td", "th"])]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            tables.append("\n".join(rows))
    if tables:
        structured["tables"] = tables

    # Headings + first following paragraph
    sections: list[str] = []
    for h in soup.find_all(["h1", "h2", "h3", "h4"]):
        heading_text = h.get_text(" ", strip=True)
        next_p = h.find_next_sibling("p")
        para_text = next_p.get_text(" ", strip=True) if next_p else ""
        if heading_text:
            sections.append(f"{heading_text}: {para_text}" if para_text else heading_text)
    if sections:
        structured["sections"] = sections

    return structured


def scrape_url(url: str, timeout: int = 12) -> list[Document]:
    """Scrape a URL and return multiple Documents (main text + structured data)."""
    if not _allowed(url):
        raise PermissionError("blocked by robots.txt")
    r = requests.get(url, timeout=timeout,
                     headers={"User-Agent": "CampusInfoBot/1.0 (student project)"})
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")

    # Extract structured data before removing tags
    structured = _extract_structured(soup)

    for t in soup(["script", "style", "nav", "footer", "header", "noscript", "form"]):
        t.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else url
    text = re.sub(r"\n{2,}", "\n", soup.get_text("\n", strip=True))

    base_meta = _meta(url, "web", categorize(text))
    base_meta["title"] = title

    docs = [Document(page_content=text[:20000], metadata=base_meta)]

    # Add table data as separate docs for better retrieval
    for table_text in structured.get("tables", []):
        if len(table_text) >= _MIN_CHUNK:
            m = _meta(f"{url} [table]", "web", categorize(table_text))
            m["title"] = title
            docs.append(Document(page_content=table_text[:5000], metadata=m))

    return docs


# ---------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------

LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".md": load_text,
    ".txt": load_text,
    ".csv": load_csv,
    ".ics": load_ics,
}


def load_all(data_dir: Path, urls: list[str]) -> tuple[list[Document], list[str]]:
    """Returns (documents, warnings). One bad source never breaks the rest."""
    docs, warnings = [], []
    for path in sorted(Path(data_dir).glob("*")):
        loader = LOADERS.get(path.suffix.lower())
        if not loader:
            continue
        try:
            new_docs = loader(path)
            docs += new_docs
        except Exception as e:
            warnings.append(f"Could not read {path.name}: {e}")
    for url in urls:
        try:
            new_docs = scrape_url(url)
            docs += new_docs
        except Exception as e:
            warnings.append(f"Could not scrape {url}: {e}")
    return docs, warnings
